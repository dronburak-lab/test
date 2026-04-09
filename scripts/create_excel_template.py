#!/usr/bin/env python3
from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
except ModuleNotFoundError:  # pragma: no cover
    Workbook = None
    DataValidation = None

try:
    from automation.task_store.base import REQUIRED_TASK_COLUMNS
except ModuleNotFoundError:  # pragma: no cover
    # Support direct execution from ./automation directory:
    # python ./scripts/create_excel_template.py
    ROOT = Path(__file__).resolve().parents[2]
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from automation.task_store.base import REQUIRED_TASK_COLUMNS


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Create Excel template for agent task pipeline")
    parser.add_argument(
        "--output",
        default="automation/tasks.xlsx",
        help="Output xlsx path (default: automation/tasks.xlsx)",
    )
    parser.add_argument(
        "--sheet-name",
        default="tasks",
        help="Worksheet name (default: tasks)",
    )
    parser.add_argument(
        "--with-example-row",
        action="store_true",
        help="Add one example task row",
    )
    return parser


def create_template(output: Path, sheet_name: str, with_example_row: bool) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    ws.append(REQUIRED_TASK_COLUMNS)
    ws.freeze_panes = "A2"

    status_values = ["planning", "ready_to_work", "in_progress", "on_review", "done"]
    priority_values = ["P0", "P1", "P2", "P3"]

    status_col = REQUIRED_TASK_COLUMNS.index("status") + 1
    priority_col = REQUIRED_TASK_COLUMNS.index("priority") + 1

    status_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(status_values)}"',
        allow_blank=False,
        showDropDown=False,
    )
    priority_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(priority_values)}"',
        allow_blank=False,
        showDropDown=False,
    )

    ws.add_data_validation(status_validation)
    ws.add_data_validation(priority_validation)

    status_validation.add(f"{ws.cell(row=1, column=status_col).column_letter}2:{ws.cell(row=1, column=status_col).column_letter}2000")
    priority_validation.add(f"{ws.cell(row=1, column=priority_col).column_letter}2:{ws.cell(row=1, column=priority_col).column_letter}2000")

    if with_example_row:
        ws.append(
            [
                "Example task title",
                "Task description",
                "planning",
                "P2",
                "",
                "",
                "",
                "",
                "",
            ]
        )

    wb.save(output)


def main() -> int:
    if Workbook is None or DataValidation is None:
        print(
            "Missing dependency: openpyxl\n"
            "Install it with: python -m pip install openpyxl",
            file=sys.stderr,
        )
        return 2

    args = build_parser().parse_args()
    create_template(Path(args.output), args.sheet_name, args.with_example_row)
    print(f"Created template: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
