from __future__ import annotations

from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from automation.status_machine import require_transition
from automation.task_store.base import (
    REQUIRED_TASK_COLUMNS,
    TaskRecord,
    TaskStatus,
    parse_task_priority,
    parse_task_status,
)

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


class ExcelTaskStore:
    def __init__(self, file_path: str, sheet_name: str = "tasks") -> None:
        self.file_path = Path(file_path)
        self.sheet_name = sheet_name
        if load_workbook is None:
            raise RuntimeError("openpyxl is required for ExcelTaskStore")

    def _load(self):
        wb = load_workbook(self.file_path)
        ws = wb[self.sheet_name]
        headers = [cell.value for cell in ws[1]]
        for required in REQUIRED_TASK_COLUMNS:
            if required not in headers:
                raise ValueError(f"Missing required column: {required}")
        return wb, ws, headers

    def _rows(self, ws, headers):
        index = {h: i for i, h in enumerate(headers)}
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[index["title"]]:
                continue
            data = {h: row[index[h]] if index[h] < len(row) else None for h in headers}
            data["status"] = parse_task_status(data.get("status"))
            data["priority"] = parse_task_priority(data.get("priority"))
            data["task_id"] = data.get("task_id") or f"row-{row_no}"
            data["meta"] = {"row_no": row_no}
            yield TaskRecord(**{k: data.get(k, "") for k in asdict(TaskRecord("", "", TaskStatus.PLANNING)).keys()})

    def list_runnable_tasks(self) -> list[TaskRecord]:
        wb, ws, headers = self._load()
        try:
            tasks = []
            for task in self._rows(ws, headers):
                if task.status == TaskStatus.READY_TO_WORK:
                    tasks.append(task)
                elif task.status == TaskStatus.ON_REVIEW and task.comment_from_user.strip():
                    tasks.append(task)
            return sorted(tasks, key=lambda x: (x.priority.value, x.task_id))
        finally:
            wb.close()

    def claim_task(self, task_id: str, expected_version: int) -> TaskRecord:
        wb, ws, headers = self._load()
        try:
            task = self._get_task(ws, headers, task_id)
            require_transition(task.status, TaskStatus.IN_PROGRESS)
            return self.update_task(task_id, {"status": TaskStatus.IN_PROGRESS.value}, expected_version)
        finally:
            wb.close()

    def append_agent_reply(self, task_id: str, text: str, expected_version: int) -> TaskRecord:
        task = self._get_by_id(task_id)
        merged = f"{task.agent_reply}\n{text}".strip()
        return self.update_task(task_id, {"agent_reply": merged}, expected_version)

    def _get_by_id(self, task_id: str) -> TaskRecord:
        wb, ws, headers = self._load()
        try:
            return self._get_task(ws, headers, task_id)
        finally:
            wb.close()

    def _get_task(self, ws, headers, task_id: str) -> TaskRecord:
        for task in self._rows(ws, headers):
            if task.task_id == task_id:
                return task
        raise KeyError(f"Task not found: {task_id}")

    def update_task(self, task_id: str, patch: dict[str, Any], expected_version: int) -> TaskRecord:
        wb, ws, headers = self._load()
        index = {h: i + 1 for i, h in enumerate(headers)}
        try:
            task = self._get_task(ws, headers, task_id)
            row_no = int(task.meta["row_no"])
            for key, value in patch.items():
                if key not in index:
                    continue
                ws.cell(row=row_no, column=index[key], value=value)
            ws.cell(row=row_no, column=index["updated_at"], value=datetime.now(timezone.utc).isoformat())
            wb.save(self.file_path)
            return self._get_task(ws, headers, task_id)
        finally:
            wb.close()
